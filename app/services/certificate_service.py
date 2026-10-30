import os
import re
import uuid
import logging
import tempfile
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from flask import current_app
from app.extensions import db
from app.models.registration import Registration
from app.models.participant import Participant
from app.models.team import Team, TeamMember
from app.models.event_setting import EventSetting
from app.models.email_log import EmailLog
from app.services.email_service import send_email

logger = logging.getLogger(__name__)

CERT_TEMPLATES_DIR = os.path.join("app", "static", "certificates", "templates")
DEFAULT_TEMPLATE_FILENAME = "synex_certificate_template.pptx"


def get_templates_dir() -> str:
    """Returns absolute path to certificate templates directory, creating it if needed."""
    full_path = os.path.abspath(CERT_TEMPLATES_DIR)
    os.makedirs(full_path, exist_ok=True)
    return full_path


def get_default_template_path() -> str:
    """Returns absolute path to the default PPTX certificate template."""
    return os.path.join(get_templates_dir(), DEFAULT_TEMPLATE_FILENAME)


def get_active_template_path() -> str:
    """Returns absolute path to current active PPTX certificate template."""
    try:
        custom_path = EventSetting.get_val("ACTIVE_CERTIFICATE_TEMPLATE", "")
        if custom_path and os.path.exists(custom_path):
            return custom_path
    except Exception:
        pass
    
    default_path = get_default_template_path()
    if not os.path.exists(default_path):
        create_default_pptx_template(default_path)
    return default_path


def create_default_pptx_template(target_path: str) -> str:
    """
    Generates an executive, professional 16:9 Landscape PPTX certificate template
    with built-in placeholder tags.
    """
    prs = Presentation()
    # 16:9 Landscape slide dimensions (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # 1. Background Shape (Dark Navy Canvas)
    bg_shape = slide.shapes.add_shape(
        1, # MSO_SHAPE.RECTANGLE
        0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
    bg_shape.line.color.rgb = RGBColor(30, 41, 59)
    
    # 2. Outer Border (Gold / Bronze Accent)
    border_shape = slide.shapes.add_shape(
        1,
        Inches(0.4), Inches(0.4),
        prs.slide_width - Inches(0.8), prs.slide_height - Inches(0.8)
    )
    border_shape.fill.background()
    border_shape.line.color.rgb = RGBColor(217, 119, 6) # Amber 600
    border_shape.line.width = Pt(3)

    # 3. Inner Decorative Border
    inner_border = slide.shapes.add_shape(
        1,
        Inches(0.52), Inches(0.52),
        prs.slide_width - Inches(1.04), prs.slide_height - Inches(1.04)
    )
    inner_border.fill.background()
    inner_border.line.color.rgb = RGBColor(51, 65, 85) # Slate 700
    inner_border.line.width = Pt(1)

    # 4. Header: Brand & Hackathon Title
    header_box = slide.shapes.add_textbox(Inches(1), Inches(0.8), Inches(11.333), Inches(1.1))
    tf_header = header_box.text_frame
    tf_header.word_wrap = True
    p_brand = tf_header.paragraphs[0]
    p_brand.alignment = PP_ALIGN.CENTER
    p_brand.text = "SYNEX NATIONAL INNOVATION HACKATHON 2026"
    p_brand.font.name = "Arial"
    p_brand.font.size = Pt(14)
    p_brand.font.bold = True
    p_brand.font.color.rgb = RGBColor(217, 119, 6) # Amber

    p_title = tf_header.add_paragraph()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.text = "CERTIFICATE OF PARTICIPATION"
    p_title.font.name = "Georgia"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # 5. Presentation line
    sub_box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(0.5))
    tf_sub = sub_box.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "PROUDLY PRESENTED TO"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(12)
    p_sub.font.bold = True
    p_sub.font.color.rgb = RGBColor(148, 163, 184) # Slate 400

    # 6. Recipient Name Box (Large Headline)
    name_box = slide.shapes.add_textbox(Inches(1), Inches(2.7), Inches(11.333), Inches(1.0))
    tf_name = name_box.text_frame
    p_name = tf_name.paragraphs[0]
    p_name.alignment = PP_ALIGN.CENTER
    p_name.text = "{{PARTICIPANT_NAME}}"
    p_name.font.name = "Georgia"
    p_name.font.size = Pt(36)
    p_name.font.bold = True
    p_name.font.color.rgb = RGBColor(56, 189, 248) # Sky 400

    # 7. Achievement Statement Paragraph
    desc_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.333), Inches(1.3))
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.alignment = PP_ALIGN.CENTER
    p_desc.text = (
        "for demonstrating exemplary innovation, teamwork, and technical excellence as {{ROLE}} "
        "in the track '{{DOMAIN}}' with squad '{{TEAM_NAME}}' representing {{COLLEGE}} "
        "at the 36-Hour National Level Hackathon."
    )
    p_desc.font.name = "Calibri"
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = RGBColor(226, 232, 240) # Slate 200

    # 8. Signatures & Metadata Footer (3 Columns)
    # Left: Issue Date & Verification ID
    left_meta = slide.shapes.add_textbox(Inches(1.2), Inches(5.4), Inches(3.5), Inches(1.2))
    tf_left = left_meta.text_frame
    p_l1 = tf_left.paragraphs[0]
    p_l1.text = "Issue Date: {{ISSUE_DATE}}"
    p_l1.font.size = Pt(11)
    p_l1.font.color.rgb = RGBColor(148, 163, 184)
    p_l2 = tf_left.add_paragraph()
    p_l2.text = "Certificate ID: {{CERTIFICATE_ID}}"
    p_l2.font.size = Pt(11)
    p_l2.font.bold = True
    p_l2.font.color.rgb = RGBColor(217, 119, 6)

    # Middle: Hackathon Official Seal
    seal_box = slide.shapes.add_textbox(Inches(5.1), Inches(5.3), Inches(3.1), Inches(1.2))
    tf_seal = seal_box.text_frame
    p_s1 = tf_seal.paragraphs[0]
    p_s1.alignment = PP_ALIGN.CENTER
    p_s1.text = "★ OFFICIAL VERIFIED ★"
    p_s1.font.size = Pt(10)
    p_s1.font.bold = True
    p_s1.font.color.rgb = RGBColor(16, 185, 129) # Emerald
    p_s2 = tf_seal.add_paragraph()
    p_s2.alignment = PP_ALIGN.CENTER
    p_s2.text = "SYNEX Academic Council"
    p_s2.font.size = Pt(10)
    p_s2.font.color.rgb = RGBColor(100, 116, 139)

    # Right: Organizing Committee Signatures
    sig_box = slide.shapes.add_textbox(Inches(8.6), Inches(5.3), Inches(3.5), Inches(1.2))
    tf_sig = sig_box.text_frame
    p_sig1 = tf_sig.paragraphs[0]
    p_sig1.alignment = PP_ALIGN.RIGHT
    p_sig1.text = "Dr. Rajeshwar Sharma"
    p_sig1.font.size = Pt(13)
    p_sig1.font.bold = True
    p_sig1.font.color.rgb = RGBColor(255, 255, 255)
    p_sig2 = tf_sig.add_paragraph()
    p_sig2.alignment = PP_ALIGN.RIGHT
    p_sig2.text = "Head of Organizing Committee & Jury Chair"
    p_sig2.font.size = Pt(10)
    p_sig2.font.color.rgb = RGBColor(148, 163, 184)

    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    prs.save(target_path)
    logger.info(f"Generated default certificate PPTX template at {target_path}")
    return target_path


def replace_placeholders_in_presentation(prs: Presentation, data: Dict[str, str]) -> None:
    """
    Replaces all placeholder strings across all slides, shapes, tables, and paragraphs
    while preserving typography, font size, bold, and color styling.
    Handles multi-run fragmented text gracefully.
    """
    replacements = {}
    for k, v in data.items():
        val_str = str(v) if v is not None else ""
        clean_k = k.strip("{}")
        replacements[f"{{{{{clean_k}}}}}"] = val_str
        replacements[f"{{{clean_k}}}"] = val_str
        replacements[f"[{clean_k}]"] = val_str
        replacements[f"%{clean_k}%"] = val_str

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                _replace_in_text_frame(shape.text_frame, replacements)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            _replace_in_text_frame(cell.text_frame, replacements)


def _replace_in_text_frame(text_frame, replacements: Dict[str, str]) -> None:
    """Performs in-place replacement on a text frame preserving style."""
    for paragraph in text_frame.paragraphs:
        full_text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
        if not full_text:
            continue

        needs_replace = False
        new_text = full_text
        for tag, replacement_val in replacements.items():
            if tag in new_text:
                new_text = new_text.replace(tag, replacement_val)
                needs_replace = True

        if needs_replace:
            if paragraph.runs:
                first_run = paragraph.runs[0]
                first_run.text = new_text
                for remaining_run in paragraph.runs[1:]:
                    remaining_run.text = ""
            else:
                paragraph.text = new_text


def convert_pptx_to_pdf(pptx_abs_path: str, pdf_abs_path: str, fallback_data: Optional[Dict[str, Any]] = None) -> bool:
    """
    Converts a filled PPTX presentation to PDF.
    Primary: Microsoft PowerPoint COM via comtypes (exact vector fidelity).
    Fallback: ReportLab PDF generator if COM is unavailable.
    """
    os.makedirs(os.path.dirname(pdf_abs_path), exist_ok=True)
    
    # 1. Attempt PowerPoint COM Automation
    try:
        import comtypes
        import comtypes.client
        
        comtypes.CoInitialize()
        try:
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.Visible = 1
            deck = powerpoint.Presentations.Open(pptx_abs_path, WithWindow=False)
            # 32 = ppSaveAsPDF
            deck.SaveAs(pdf_abs_path, 32)
            deck.Close()
            powerpoint.Quit()
            if os.path.exists(pdf_abs_path) and os.path.getsize(pdf_abs_path) > 1000:
                logger.info(f"Converted PPTX to PDF via PowerPoint COM: {pdf_abs_path}")
                return True
        finally:
            comtypes.CoUninitialize()
    except Exception as com_err:
        logger.warning(f"PowerPoint COM conversion failed or unavailable ({com_err}). Falling back to ReportLab...")

    # 2. Fallback: ReportLab High-Resolution Vector PDF
    try:
        generate_reportlab_certificate_pdf(fallback_data or {}, pdf_abs_path)
        if os.path.exists(pdf_abs_path) and os.path.getsize(pdf_abs_path) > 500:
            logger.info(f"Generated certificate PDF via ReportLab fallback: {pdf_abs_path}")
            return True
    except Exception as rl_err:
        logger.error(f"ReportLab certificate generation failed: {rl_err}")
    
    return False


def generate_reportlab_certificate_pdf(data: Dict[str, Any], output_pdf_path: str) -> str:
    """Generates an executive-style landscape Certificate PDF using ReportLab."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor

    p_name = data.get("PARTICIPANT_NAME") or data.get("name") or "Participant"
    role = data.get("ROLE") or data.get("role") or "Participant"
    domain = data.get("DOMAIN") or data.get("domain") or "Innovation Track"
    team = data.get("TEAM_NAME") or data.get("team_name") or "Individual"
    college = data.get("COLLEGE") or data.get("college") or "Recognized Institution"
    cert_id = data.get("CERTIFICATE_ID") or data.get("certificate_id") or "SYNEX-CERT"
    issue_date = data.get("ISSUE_DATE") or data.get("date") or datetime.utcnow().strftime("%B %d, %Y")

    width, height = landscape(letter) # 792 x 612 pt
    c = canvas.Canvas(output_pdf_path, pagesize=(width, height))

    # Dark background
    c.setFillColor(HexColor("#0F172A"))
    c.rect(0, 0, width, height, fill=1, stroke=0)

    # Gold outer frame
    c.setStrokeColor(HexColor("#D97706"))
    c.setLineWidth(4)
    c.rect(30, 30, width - 60, height - 60, fill=0, stroke=1)

    # Inner subtle frame
    c.setStrokeColor(HexColor("#334155"))
    c.setLineWidth(1)
    c.rect(42, 42, width - 84, height - 84, fill=0, stroke=1)

    # Header Top Brand
    c.setFillColor(HexColor("#D97706"))
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(width / 2.0, height - 80, "SYNEX NATIONAL INNOVATION HACKATHON 2026")

    # Title
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Times-Bold", 28)
    c.drawCentredString(width / 2.0, height - 125, "CERTIFICATE OF PARTICIPATION")

    # Subtitle
    c.setFillColor(HexColor("#94A3B8"))
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2.0, height - 165, "PROUDLY PRESENTED TO")

    # Recipient Name
    c.setFillColor(HexColor("#38BDF8"))
    c.setFont("Times-BoldItalic", 34)
    c.drawCentredString(width / 2.0, height - 215, str(p_name))

    # Dividing accent line
    c.setStrokeColor(HexColor("#D97706"))
    c.setLineWidth(1.5)
    c.line(width / 2.0 - 180, height - 230, width / 2.0 + 180, height - 230)

    # Citation / Achievement text
    c.setFillColor(HexColor("#E2E8F0"))
    c.setFont("Helvetica", 13)
    c.drawCentredString(width / 2.0, height - 265, f"for demonstrating exemplary innovation and technical excellence as {role}")
    c.drawCentredString(width / 2.0, height - 288, f"in the '{domain}' track with squad '{team}'")
    c.drawCentredString(width / 2.0, height - 311, f"representing {college} during the 36-hour national hackathon.")

    # Bottom Left: Date & ID
    c.setFillColor(HexColor("#94A3B8"))
    c.setFont("Helvetica", 10)
    c.drawString(65, 120, f"Issue Date: {issue_date}")
    c.setFillColor(HexColor("#D97706"))
    c.setFont("Helvetica-Bold", 10)
    c.drawString(65, 102, f"Certificate ID: {cert_id}")

    # Center Verification Seal
    c.setFillColor(HexColor("#10B981"))
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(width / 2.0, 115, "★ OFFICIAL VERIFIED PARTICIPANT ★")
    c.setFillColor(HexColor("#64748B"))
    c.setFont("Helvetica", 9)
    c.drawCentredString(width / 2.0, 98, "SYNEX Academic Steering Committee")

    # Right: Signature Block
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica-Bold", 12)
    c.drawRightString(width - 65, 120, "Dr. Rajeshwar Sharma")
    c.setFillColor(HexColor("#94A3B8"))
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 65, 102, "Head of Organizing Committee & Jury Chair")

    c.save()
    return output_pdf_path


def get_participants_for_certificates(
    domain_filter: str = "all",
    paid_only: bool = True,
    search: str = ""
) -> List[Dict[str, Any]]:
    """
    Extracts flat list of all eligible participants (Solo Hackers, Team Leaders, Team Members)
    with their metadata prepared for certificate generation.
    """
    query = Registration.query
    if paid_only:
        query = query.filter(Registration.payment_status == "paid")
    
    if domain_filter and domain_filter.lower() != "all":
        domain_term = f"%{domain_filter.strip()}%"
        from sqlalchemy import or_
        query = query.filter(or_(Registration.domain == domain_filter, Registration.domain.ilike(domain_term)))

    registrations = query.order_by(Registration.created_at.desc()).all()
    participants_list = []
    
    for reg in registrations:
        is_team = reg.registration_type == "team"
        lead = reg.participant or Participant(full_name="N/A", email="", phone="", college="")
        team_title = reg.team.team_name if reg.team else "Individual Participant"
        
        # 1. Team Leader or Solo Hacker
        lead_role = "Team Leader" if is_team else "Solo Hacker"
        lead_id = f"reg_{reg.id}_lead"
        cert_id = f"SYNEX-CERT-{reg.public_id}-01"
        
        participants_list.append({
            "id": lead_id,
            "registration_id": reg.id,
            "public_id": reg.public_id,
            "name": lead.full_name or "N/A",
            "email": lead.email or "",
            "phone": lead.phone or "",
            "role": lead_role,
            "domain": reg.domain or "AI & Machine Learning",
            "team_name": team_title,
            "college": lead.college or "N/A",
            "payment_status": reg.payment_status or "pending",
            "certificate_id": cert_id,
            "issue_date": datetime.utcnow().strftime("%B %d, %Y"),
        })

        # 2. Additional Team Members
        if is_team and reg.team and reg.team.members:
            for idx, member in enumerate(reg.team.members, start=2):
                mem_id = f"reg_{reg.id}_mem_{member.id}"
                mem_cert_id = f"SYNEX-CERT-{reg.public_id}-{idx:02d}"
                participants_list.append({
                    "id": mem_id,
                    "registration_id": reg.id,
                    "public_id": reg.public_id,
                    "name": member.name or "N/A",
                    "email": member.email or "",
                    "phone": member.phone or "",
                    "role": "Team Member",
                    "domain": reg.domain or "AI & Machine Learning",
                    "team_name": team_title,
                    "college": member.college or lead.college or "N/A",
                    "payment_status": reg.payment_status or "pending",
                    "certificate_id": mem_cert_id,
                    "issue_date": datetime.utcnow().strftime("%B %d, %Y"),
                })

    if search:
        s = search.strip().lower()
        participants_list = [
            p for p in participants_list
            if s in p["name"].lower()
            or s in p["email"].lower()
            or s in p["domain"].lower()
            or s in p["team_name"].lower()
            or s in p["college"].lower()
            or s in p["public_id"].lower()
        ]

    return participants_list


def generate_single_certificate_pdf(participant_data: Dict[str, Any]) -> Tuple[bytes, str]:
    """
    Fills active PPTX template for a participant and converts to PDF.
    Returns (pdf_bytes, filename).
    """
    active_template = get_active_template_path()
    p_name = participant_data.get("name") or "Participant"
    clean_name = re.sub(r"[^a-zA-Z0-9_-]", "_", p_name)
    public_id = participant_data.get("public_id") or "SYNEX"
    pdf_filename = f"SYNEX_Certificate_{public_id}_{clean_name}.pdf"

    context = {
        "PARTICIPANT_NAME": p_name,
        "NAME": p_name,
        "ROLE": participant_data.get("role") or "Participant",
        "DOMAIN": participant_data.get("domain") or "Innovation Track",
        "TRACK": participant_data.get("domain") or "Innovation Track",
        "TEAM_NAME": participant_data.get("team_name") or "Individual",
        "COLLEGE": participant_data.get("college") or "Recognized Institution",
        "INSTITUTION": participant_data.get("college") or "Recognized Institution",
        "CERTIFICATE_ID": participant_data.get("certificate_id") or f"SYNEX-CERT-{public_id}",
        "REF_ID": public_id,
        "ISSUE_DATE": participant_data.get("issue_date") or datetime.utcnow().strftime("%B %d, %Y"),
        "DATE": participant_data.get("issue_date") or datetime.utcnow().strftime("%B %d, %Y"),
        "EVENT_NAME": "SYNEX National Innovation Hackathon 2026",
    }

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_pptx = os.path.join(tmpdir, "cert_filled.pptx")
        tmp_pdf = os.path.join(tmpdir, pdf_filename)

        prs = Presentation(active_template)
        replace_placeholders_in_presentation(prs, context)
        prs.save(tmp_pptx)

        success = convert_pptx_to_pdf(tmp_pptx, tmp_pdf, fallback_data=context)
        if not success or not os.path.exists(tmp_pdf):
            generate_reportlab_certificate_pdf(context, tmp_pdf)

        with open(tmp_pdf, "rb") as f:
            pdf_bytes = f.read()

    return pdf_bytes, pdf_filename


def send_participant_certificate_email(
    participant_data: Dict[str, Any],
    custom_subject: Optional[str] = None,
    custom_body: Optional[str] = None
) -> Tuple[bool, str, Optional[EmailLog]]:
    """
    Generates personalized certificate PDF and dispatches it via email to participant.
    """
    to_email = (participant_data.get("email") or "").strip()
    if not to_email:
        return False, "Participant does not have an email address on file.", None

    p_name = participant_data.get("name") or "Participant"
    domain = participant_data.get("domain") or "Innovation Track"
    cert_id = participant_data.get("certificate_id") or "SYNEX-CERT"

    subject = custom_subject or f"Your SYNEX 2026 Certificate of Participation — {p_name}"
    
    pdf_bytes, pdf_filename = generate_single_certificate_pdf(participant_data)

    body_text = custom_body or (
        f"Dear {p_name},\n\n"
        f"Congratulations on your outstanding contribution at the SYNEX National Innovation Hackathon 2026!\n\n"
        f"We are delighted to present your official Certificate of Participation for the '{domain}' track.\n"
        f"Your Certificate Verification ID is: {cert_id}.\n\n"
        f"Please find your high-resolution official Certificate attached as a PDF document.\n\n"
        f"Warm regards,\n"
        f"SYNEX Organizing Committee\n"
        f"Campus Innovation Hub, Bengaluru\n"
        f"https://synex2026.org"
    )

    body_html = f"""
    <div style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; max-width: 600px; margin: 0 auto; padding: 28px; border: 1px solid #E2E8F0; border-radius: 12px; background: #FFFFFF;">
      <div style="padding-bottom: 18px; margin-bottom: 22px; border-bottom: 3px solid #2563EB;">
        <span style="font-size: 11px; text-transform: uppercase; font-weight: 800; color: #D97706; letter-spacing: 1px;">Official Credentials</span>
        <h2 style="margin: 6px 0 0; color: #1E293B; font-size: 22px; font-weight: 800;">SYNEX 2026 Participation Certificate</h2>
      </div>
      <p style="color: #334155; font-size: 15px; line-height: 1.6;">Dear <strong>{p_name}</strong>,</p>
      <p style="color: #334155; font-size: 15px; line-height: 1.6;">
        Congratulations on your active participation and technical contributions during the <strong>SYNEX National Innovation Hackathon 2026</strong>!
      </p>
      <div style="background: #F8FAFC; border: 1px solid #E2E8F0; border-left: 4px solid #10B981; border-radius: 6px; padding: 14px 16px; margin: 20px 0;">
        <div style="font-size: 13px; color: #64748B;">Innovation Track: <strong style="color: #0F172A;">{domain}</strong></div>
        <div style="font-size: 13px; color: #64748B; margin-top: 4px;">Role: <strong style="color: #0F172A;">{participant_data.get('role', 'Participant')}</strong></div>
        <div style="font-size: 13px; color: #64748B; margin-top: 4px;">Certificate Verification ID: <strong style="color: #2563EB; font-family: monospace;">{cert_id}</strong></div>
      </div>
      <p style="color: #334155; font-size: 15px; line-height: 1.6;">
        Your verified e-Certificate is attached to this email as a PDF document. You may also showcase this achievement on LinkedIn and technical portfolios.
      </p>
      <div style="margin-top: 32px; padding-top: 18px; border-top: 1px solid #E2E8F0; font-size: 12px; color: #94A3B8;">
        <p style="margin: 0;"><strong>SYNEX 2026 Organizing Committee</strong> &bull; Campus Innovation Hub, Bengaluru</p>
        <p style="margin: 4px 0 0;">Visit <a href="https://synex2026.org" style="color: #2563EB;">synex2026.org</a> for winner announcements and post-hackathon highlights.</p>
      </div>
    </div>
    """

    attachments = [{
        "filename": pdf_filename,
        "content_bytes": pdf_bytes,
        "content_type": "application/pdf"
    }]

    success, msg, log = send_email(
        to_email=to_email,
        subject=subject,
        body_text=body_text,
        body_html=body_html,
        email_type="certificate",
        attachments=attachments,
    )

    return success, msg, log
